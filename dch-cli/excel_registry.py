"""Read and write the DCH Excel sheet (dch/dch_registry.xlsx).

Owner: Teammate 4. Used by update_dch.py, create_registry.py and reset_demo.py.

The workbook has two sheets:
  Assets  one row per dataset, keyed by urn
  Audit   one row per create, update or validation failure, never edited or deleted

Columns are found by their header text, so reordering columns is safe.
Renaming or deleting a header is not.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ASSET_SHEET = "Assets"
AUDIT_SHEET = "Audit"

ASSET_HEADERS = [
    "urn", "dataset", "data_product_name", "owner", "domain", "description",
    "location", "format", "file_count", "total_bytes", "first_seen",
    "last_updated", "last_run_id",
]
AUDIT_HEADERS = ["timestamp", "urn", "action", "changed_fields", "run_id", "commit_sha"]

# Fields compared on every run. A difference in any of them updates the row.
TRACKED_FIELDS = [
    "data_product_name", "owner", "domain", "description",
    "file_count", "total_bytes", "last_updated",
]

COLUMN_WIDTHS = {
    "urn": 64, "dataset": 16, "data_product_name": 24, "owner": 28, "domain": 14,
    "description": 48, "location": 52, "format": 9, "file_count": 11,
    "total_bytes": 13, "first_seen": 22, "last_updated": 22, "last_run_id": 14,
    "timestamp": 22, "action": 10, "changed_fields": 80, "run_id": 14, "commit_sha": 42,
}


def create_workbook(path: str | Path) -> None:
    """Write a blank workbook with styled headers, a frozen header row and set widths."""
    wb = Workbook()
    assets = wb.active
    assets.title = ASSET_SHEET
    audit = wb.create_sheet(AUDIT_SHEET)

    for ws, headers in ((assets, ASSET_HEADERS), (audit, AUDIT_HEADERS)):
        ws.append(headers)
        for col, name in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E4C74")
            ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS.get(name, 16)
        ws.freeze_panes = "A2"

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _same(old, new) -> bool:
    """Compare a cell value with a new value, treating 3 and 3.0 as equal."""
    if isinstance(old, (int, float)) and isinstance(new, (int, float)):
        return float(old) == float(new)
    return ("" if old is None else str(old)) == ("" if new is None else str(new))


class ExcelRegistry:
    """The DCH sheet loaded into memory. Call upsert() once per dataset, then save_if_changed()."""

    def __init__(self, path: str | Path, run_id: str, commit_sha: str, now: str):
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(
                f"{self.path} not found. Create it with: python dch-cli/create_registry.py"
            )
        self.run_id = run_id
        self.commit_sha = commit_sha
        self.now = now

        self.wb = load_workbook(self.path)
        self.assets = self._sheet(ASSET_SHEET)
        self.audit = self._sheet(AUDIT_SHEET)
        self.asset_cols = self._columns(self.assets, ASSET_HEADERS)
        self.audit_cols = self._columns(self.audit, AUDIT_HEADERS)

        self.rows = self._index_rows()  # urn -> row number
        self.next_asset_row = max(self.rows.values(), default=1) + 1
        self.next_audit_row = self._last_filled_row(self.audit, self.audit_cols["timestamp"]) + 1
        self.change_count = 0

    # ---------- loading ----------

    def _sheet(self, name: str):
        if name not in self.wb.sheetnames:
            raise ValueError(f"{self.path} has no '{name}' sheet. Sheets found: {self.wb.sheetnames}")
        return self.wb[name]

    @staticmethod
    def _columns(ws, required: list[str]) -> dict[str, int]:
        found = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
        missing = [h for h in required if h not in found]
        if missing:
            raise ValueError(f"Sheet '{ws.title}' is missing header(s): {', '.join(missing)}")
        return found

    def _index_rows(self) -> dict[str, int]:
        col = self.asset_cols["urn"]
        index: dict[str, int] = {}
        for row in range(2, self.assets.max_row + 1):
            urn = self.assets.cell(row=row, column=col).value
            if urn:
                index[str(urn).strip()] = row
        return index

    @staticmethod
    def _last_filled_row(ws, col: int) -> int:
        last = 1
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=col).value not in (None, ""):
                last = row
        return last

    # ---------- reading ----------

    def get(self, urn: str) -> dict | None:
        row = self.rows.get(urn)
        if row is None:
            return None
        return {h: self.assets.cell(row=row, column=c).value for h, c in self.asset_cols.items()}

    def audit_rows(self) -> list[dict]:
        rows = []
        for row in range(2, self.next_audit_row):
            rows.append({h: self.audit.cell(row=row, column=c).value for h, c in self.audit_cols.items()})
        return rows

    # ---------- writing ----------

    def upsert(self, record: dict) -> str:
        """Create or update one dataset row. Returns 'created', 'updated' or 'unchanged'.

        record needs: urn, dataset, data_product_name, owner, domain, description,
        location, format, file_count, total_bytes, last_updated
        """
        urn = record["urn"]
        row = self.rows.get(urn)

        if row is None:
            row = self.next_asset_row
            self.next_asset_row += 1
            values = dict(record, first_seen=self.now, last_run_id=self.run_id)
            for header in ASSET_HEADERS:
                self.assets.cell(row=row, column=self.asset_cols[header], value=values.get(header))
            self.rows[urn] = row
            self._audit(urn, "created", {f: {"old": None, "new": values[f]} for f in TRACKED_FIELDS})
            return "created"

        diff = {}
        for f in TRACKED_FIELDS:
            old = self.assets.cell(row=row, column=self.asset_cols[f]).value
            if not _same(old, record[f]):
                diff[f] = {"old": old, "new": record[f]}
        if not diff:
            return "unchanged"

        for f, change in diff.items():
            self.assets.cell(row=row, column=self.asset_cols[f], value=change["new"])
        self.assets.cell(row=row, column=self.asset_cols["last_run_id"], value=self.run_id)
        self._audit(urn, "updated", diff)
        return "updated"

    def log_validation_failure(self, urn: str, missing: list[str]) -> None:
        """Log a dataset whose metadata is incomplete. Never touches Assets.

        Skipped when the latest Audit row for this urn already logs the same failure,
        so reruns stay no-ops.
        """
        fields = {"missing": missing}
        for row in reversed(self.audit_rows()):
            if row["urn"] == urn:
                if row["action"] == "validation_failed" and row["changed_fields"] == json.dumps(fields, default=str):
                    return
                break
        self._audit(urn, "validation_failed", fields)

    def _audit(self, urn: str, action: str, fields: dict) -> None:
        values = {
            "timestamp": self.now,
            "urn": urn,
            "action": action,
            "changed_fields": json.dumps(fields, default=str),
            "run_id": self.run_id,
            "commit_sha": self.commit_sha,
        }
        for header, value in values.items():
            self.audit.cell(row=self.next_audit_row, column=self.audit_cols[header], value=value)
        self.next_audit_row += 1
        self.change_count += 1

    def save_if_changed(self) -> bool:
        """Save only when something changed.

        openpyxl writes different bytes on every save, even when the content is the
        same, so saving on every run would make Git see a change every time.
        """
        if self.change_count == 0:
            return False
        self.wb.save(self.path)
        return True
