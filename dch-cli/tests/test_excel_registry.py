"""Unit tests for excel_registry.py.

Owner: Teammate 4. Run from the repo root:  pytest dch-cli/tests -q
"""
import json
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # make dch-cli importable
from excel_registry import (  # noqa: E402
    ASSET_HEADERS, AUDIT_HEADERS, ExcelRegistry, create_workbook,
)

URN = "urn:dch:dataset:databricks:dev:workspace.dch_demo.landing/customers"


def record(file_count=1, total_bytes=100, last_updated="2026-09-21T10:00:00Z", owner="data-team@example.com"):
    return {
        "urn": URN,
        "dataset": "customers",
        "data_product_name": "Customer Master",
        "owner": owner,
        "domain": "sales",
        "description": "Customer records landed from the CRM export.",
        "location": "/Volumes/workspace/dch_demo/landing/customers/",
        "format": "json",
        "file_count": file_count,
        "total_bytes": total_bytes,
        "last_updated": last_updated,
    }


def open_registry(path, run_id="1"):
    return ExcelRegistry(path, run_id=run_id, commit_sha="abc123", now="2026-09-21T12:00:00Z")


@pytest.fixture
def sheet(tmp_path):
    path = tmp_path / "dch_registry.xlsx"
    create_workbook(path)
    return path


def test_create_workbook_has_both_sheets_and_headers(sheet):
    wb = load_workbook(sheet)
    assert wb.sheetnames == ["Assets", "Audit"]
    assert [c.value for c in wb["Assets"][1]] == ASSET_HEADERS
    assert [c.value for c in wb["Audit"][1]] == AUDIT_HEADERS
    assert wb["Assets"].freeze_panes == "A2"


def test_new_urn_appends_row_and_audit(sheet):
    reg = open_registry(sheet)
    assert reg.upsert(record()) == "created"
    assert reg.save_if_changed() is True

    reg = open_registry(sheet)
    row = reg.get(URN)
    assert row["file_count"] == 1
    assert row["first_seen"] == "2026-09-21T12:00:00Z"
    audit = reg.audit_rows()
    assert [a["action"] for a in audit] == ["created"]


def test_changed_fields_update_row_and_log_old_and_new(sheet):
    reg = open_registry(sheet)
    reg.upsert(record())
    reg.save_if_changed()

    reg = open_registry(sheet, run_id="2")
    assert reg.upsert(record(file_count=2, total_bytes=250)) == "updated"
    reg.save_if_changed()

    reg = open_registry(sheet)
    assert len(reg.rows) == 1
    assert reg.get(URN)["last_run_id"] == "2"
    last = reg.audit_rows()[-1]
    assert last["action"] == "updated"
    changed = json.loads(last["changed_fields"])
    assert changed == {"file_count": {"old": 1, "new": 2}, "total_bytes": {"old": 100, "new": 250}}


def test_same_values_are_unchanged_and_file_is_not_saved(sheet):
    reg = open_registry(sheet)
    reg.upsert(record())
    reg.save_if_changed()
    before = sheet.read_bytes()

    reg = open_registry(sheet)
    assert reg.upsert(record()) == "unchanged"
    assert reg.save_if_changed() is False
    assert sheet.read_bytes() == before


def test_columns_are_found_by_header_name(tmp_path):
    path = tmp_path / "shuffled.xlsx"
    wb = Workbook()
    assets = wb.active
    assets.title = "Assets"
    assets.append(list(reversed(ASSET_HEADERS)))
    wb.create_sheet("Audit").append(AUDIT_HEADERS)
    wb.save(path)

    reg = open_registry(path)
    reg.upsert(record(file_count=7))
    reg.save_if_changed()

    ws = load_workbook(path)["Assets"]
    headers = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=headers.index("file_count") + 1).value == 7
    assert ws.cell(row=2, column=headers.index("urn") + 1).value == URN


def test_missing_header_gives_clear_error(tmp_path):
    path = tmp_path / "broken.xlsx"
    wb = Workbook()
    wb.active.title = "Assets"
    wb.active.append([h for h in ASSET_HEADERS if h != "total_bytes"])
    wb.create_sheet("Audit").append(AUDIT_HEADERS)
    wb.save(path)

    with pytest.raises(ValueError, match="total_bytes"):
        open_registry(path)


def test_missing_workbook_says_how_to_create_it(tmp_path):
    with pytest.raises(FileNotFoundError, match="create_registry.py"):
        open_registry(tmp_path / "nope.xlsx")


def test_metadata_change_updates_existing_row(sheet):
    reg = open_registry(sheet)
    reg.upsert(record())
    reg.save_if_changed()

    reg = open_registry(sheet, run_id="2")
    assert reg.upsert(record(owner="new-owner@example.com")) == "updated"
    reg.save_if_changed()

    reg = open_registry(sheet)
    assert len(reg.rows) == 1
    assert reg.get(URN)["owner"] == "new-owner@example.com"
    changed = json.loads(reg.audit_rows()[-1]["changed_fields"])
    assert changed == {"owner": {"old": "data-team@example.com", "new": "new-owner@example.com"}}


def test_validation_failure_is_logged_once(sheet):
    reg = open_registry(sheet)
    reg.log_validation_failure(URN, ["owner", "domain"])
    assert reg.save_if_changed() is True

    reg = open_registry(sheet, run_id="2")
    reg.log_validation_failure(URN, ["owner", "domain"])
    assert reg.save_if_changed() is False

    reg = open_registry(sheet)
    audit = reg.audit_rows()
    assert [a["action"] for a in audit] == ["validation_failed"]
    assert json.loads(audit[0]["changed_fields"]) == {"missing": ["owner", "domain"]}
    assert reg.rows == {}
